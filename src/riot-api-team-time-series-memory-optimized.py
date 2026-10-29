# -*- coding: utf-8 -*-
"""Extrai partidas da Riot API e gera diretamente a planilha de série temporal.

A saída mantém a estrutura do arquivo team-time-series.xlsx de referência:
aba Series_Temporais, duas linhas por minuto (teamId 100 e 200), vencedor,
variáveis agregadas por equipe, eventos acumulados e diferenças azul - vermelho.
Não cria uma planilha intermediária de dados brutos.
"""

from __future__ import annotations

import json
import math
import time
from collections import defaultdict, deque
from urllib.parse import urlparse

import requests
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


# ============================================================
# 1. CONFIGURAÇÕES — EDITE ESTA SEÇÃO
# ============================================================

RIOT_API_KEY = "COLE_SUA_CHAVE_DA_RIOT_API_AQUI"  # Cole sua chave da Riot API aqui

quantidade_jogadores = 200       # N melhores jogadores do ranking BR
quantidade_partidas = 50        # M partidas válidas e únicas por jogador
duracao_minima_segundos = 20 * 60
plataforma = "br1"
regiao_match = "americas"
queue_id = 420                  # Ranqueada Solo/Duo
map_id = 11
game_mode = "CLASSIC"

nome_arquivo_saida = "../data/team-time-series.xlsx"

# Limites de referência da aplicação. O limitador permite rajadas curtas,
# mas respeita as duas janelas. Se a Riot retornar HTTP 429, obedece Retry-After.
LIMITE_CURTO = 20
JANELA_CURTA = 1.0
LIMITE_LONGO = 100
JANELA_LONGA = 120.0
MARGEM_SEGURANCA = 0.05
TIMEOUT_REQUISICAO = 30

EQUIPES = (100, 200)

COLUNAS_SOMA = [
    "championStats.abilityHaste",
    "championStats.abilityPower",
    "championStats.armor",
    "championStats.attackDamage",
    "championStats.health",
    "currentGold",
    "damageStats.magicDamageDone",
    "damageStats.magicDamageDoneToChampions",
    "damageStats.magicDamageTaken",
    "damageStats.physicalDamageDone",
    "damageStats.physicalDamageDoneToChampions",
    "damageStats.physicalDamageTaken",
    "damageStats.totalDamageDone",
    "damageStats.totalDamageDoneToChampions",
    "damageStats.totalDamageTaken",
    "damageStats.trueDamageDone",
    "damageStats.trueDamageDoneToChampions",
    "damageStats.trueDamageTaken",
    "goldPerSecond",
    "jungleMinionsKilled",
    "level",
    "minionsKilled",
    "timeEnemySpentControlled",
    "totalGold",
    "xp",
]

COLUNAS_EVENTOS = [
    "kills", "deaths", "assists", "dragons", "elder_dragons", "barons",
    "rift_heralds", "atakhan", "towers_destroyed", "inhibitors_destroyed",
]


# ============================================================
# 2. SESSÃO HTTP E LIMITADOR DE REQUISIÇÕES
# ============================================================

if not RIOT_API_KEY or RIOT_API_KEY == "COLE_SUA_CHAVE_DA_RIOT_API_AQUI":
    raise SystemExit(
        "Edite RIOT_API_KEY no início do arquivo e cole sua chave válida da Riot API."
    )

sessao = requests.Session()
sessao.headers.update({"X-Riot-Token": RIOT_API_KEY})
# Cada host tem seu próprio histórico local de requisições.
historico_limites = defaultdict(lambda: {"curto": deque(), "longo": deque()})


def aguardar_limite(host: str) -> None:
    historico = historico_limites[host]
    while True:
        agora = time.monotonic()
        while historico["curto"] and agora - historico["curto"][0] >= JANELA_CURTA:
            historico["curto"].popleft()
        while historico["longo"] and agora - historico["longo"][0] >= JANELA_LONGA:
            historico["longo"].popleft()

        esperas = [0.0]
        if len(historico["curto"]) >= LIMITE_CURTO:
            esperas.append(historico["curto"][0] + JANELA_CURTA + MARGEM_SEGURANCA - agora)
        if len(historico["longo"]) >= LIMITE_LONGO:
            esperas.append(historico["longo"][0] + JANELA_LONGA + MARGEM_SEGURANCA - agora)

        espera = max(esperas)
        if espera <= 0:
            historico["curto"].append(agora)
            historico["longo"].append(agora)
            return
        print(f"[LIMITE LOCAL] Aguardando {espera:.1f}s para respeitar os limites da API...")
        time.sleep(espera)


def tempo_retry_after(resposta: requests.Response) -> float:
    valor = resposta.headers.get("Retry-After")
    if valor:
        try:
            return max(1.0, float(valor))
        except ValueError:
            pass
    return 10.0


def requisitar_json(url: str, params: dict | None = None, tentativas_conexao: int = 4):
    """Faz GET com limitador, retry para erros transitórios e espera em HTTP 429."""
    host = urlparse(url).netloc
    falhas_conexao = 0
    while True:
        aguardar_limite(host)
        try:
            resposta = sessao.get(url, params=params, timeout=TIMEOUT_REQUISICAO)
        except requests.RequestException as erro:
            falhas_conexao += 1
            if falhas_conexao > tentativas_conexao:
                print(f"[ERRO DE CONEXÃO] {url}: {erro}")
                return None, None
            espera = min(2 ** falhas_conexao, 20)
            print(f"[CONEXÃO] Falha temporária; nova tentativa em {espera}s ({falhas_conexao}/{tentativas_conexao}).")
            time.sleep(espera)
            continue

        if resposta.status_code == 429:
            espera = tempo_retry_after(resposta)
            tipo = resposta.headers.get("X-Rate-Limit-Type", "não informado")
            resposta.close()
            print(f"[HTTP 429] Limite Riot ({tipo}); aguardando {espera:.1f}s e retomando.")
            time.sleep(espera)
            continue

        if resposta.status_code in (500, 502, 503, 504):
            falhas_conexao += 1
            status = resposta.status_code
            resposta.close()
            if falhas_conexao > tentativas_conexao:
                print(f"[ERRO HTTP {status}] Número máximo de tentativas excedido: {url}")
                return status, None
            espera = min(2 ** falhas_conexao, 20)
            print(f"[HTTP {status}] Erro temporário; nova tentativa em {espera}s.")
            time.sleep(espera)
            continue

        try:
            dados = resposta.json() if resposta.content else None
        except ValueError:
            dados = None
        status = resposta.status_code
        if status >= 400:
            detalhe = resposta.text[:250].replace("\n", " ")
            print(f"[HTTP {status}] {url} — {detalhe}")
        resposta.close()
        return status, dados


# ============================================================
# 3. SELEÇÃO DOS N MELHORES JOGADORES BRASILEIROS
# ============================================================

def selecionar_jogadores() -> list[dict]:
    print("\n" + "=" * 72)
    print("ETAPA 1/4 — CONSULTA AO RANKING BRASILEIRO")
    print("=" * 72)
    base = f"https://{plataforma}.api.riotgames.com"
    tiers = [
        ("CHALLENGER", "challengerleagues"),
        ("GRANDMASTER", "grandmasterleagues"),
        ("MASTER", "masterleagues"),
    ]
    ordem_tier = {"CHALLENGER": 0, "GRANDMASTER": 1, "MASTER": 2}
    candidatos = []

    for tier, endpoint in tiers:
        url = f"{base}/lol/league/v4/{endpoint}/by-queue/RANKED_SOLO_5x5"
        status, dados = requisitar_json(url)
        if status != 200 or not isinstance(dados, dict):
            if status == 403:
                raise RuntimeError("HTTP 403 no ranking. Verifique a chave e suas permissões.")
            continue
        entradas = dados.get("entries", [])
        for jogador in entradas:
            jogador["tier"] = tier
            candidatos.append(jogador)
        print(f"{tier}: {len(entradas)} jogadores no ranking.")

    candidatos.sort(key=lambda j: (ordem_tier[j["tier"]], -j.get("leaguePoints", 0)))
    selecionados, vistos = [], set()
    print(f"Candidatos encontrados: {len(candidatos)}. Recuperando PUUIDs somente quando necessário...")

    for candidato in candidatos:
        if len(selecionados) >= quantidade_jogadores:
            break
        summoner_id = candidato.get("summonerId")
        puuid = candidato.get("puuid")
        identificador = puuid or summoner_id
        if not identificador or identificador in vistos:
            continue
        vistos.add(identificador)

        # League-V4 normalmente fornece summonerId, mas não PUUID.
        if not puuid and summoner_id:
            url = f"{base}/lol/summoner/v4/summoners/{summoner_id}"
            status, dados = requisitar_json(url)
            if status != 200 or not isinstance(dados, dict):
                continue
            puuid = dados.get("puuid")
        if not puuid:
            continue

        jogador = {
            "puuid": puuid,
            "nome": f"Jogador {len(selecionados) + 1}",
            "tier": candidato["tier"],
            "league_points": candidato.get("leaguePoints", 0),
        }
        selecionados.append(jogador)
        print(f"Selecionado {len(selecionados)}/{quantidade_jogadores}: {jogador['tier']} | {jogador['league_points']} LP | PUUID {puuid[:12]}…")

    if not selecionados:
        raise RuntimeError("Nenhum jogador foi selecionado. Verifique a chave e os endpoints do ranking.")
    print(f"\nJogadores selecionados: {len(selecionados)}")
    return selecionados


# ============================================================
# 4. BUSCA DE PARTIDAS VÁLIDAS E ÚNICAS
# ============================================================

def coletar_partidas(jogadores: list[dict]) -> dict[str, dict]:
    print("\n" + "=" * 72)
    print("ETAPA 2/4 — EXTRAÇÃO DAS PARTIDAS")
    print("=" * 72)
    partidas_unicas: dict[str, dict] = {}
    dono_partida: dict[str, str] = {}
    cache_detalhes: dict[str, dict | None] = {}
    conflitos = 0
    partidas_examindas = 0

    for indice, jogador in enumerate(jogadores, start=1):
        puuid = jogador["puuid"]
        nome = jogador["nome"]
        validas_jogador = 0
        analisadas_jogador = 0
        conflitos_jogador = 0
        inicio = 0
        tamanho_pagina = 100
        print("\n" + "-" * 72)
        print(f"Jogador {indice}/{len(jogadores)} — {nome} | {jogador['tier']} | {jogador['league_points']} LP")

        while validas_jogador < quantidade_partidas:
            url_historico = f"https://{regiao_match}.api.riotgames.com/lol/match/v5/matches/by-puuid/{puuid}/ids"
            status, ids = requisitar_json(url_historico, params={"start": inicio, "count": tamanho_pagina, "queue": queue_id})
            if status != 200 or not isinstance(ids, list):
                print(f"[HISTÓRICO INDISPONÍVEL] {nome}; interrompendo a busca deste jogador.")
                break
            if not ids:
                print(f"[FIM DO HISTÓRICO] Não há mais partidas para {nome}.")
                break
            inicio += len(ids)

            for match_id in ids:
                if validas_jogador >= quantidade_partidas:
                    break
                analisadas_jogador += 1
                partidas_examindas += 1

                # Se já temos detalhes, reutilizamos: evita chamadas repetidas quando
                # a mesma partida aparece no histórico de mais de um jogador.
                if match_id not in cache_detalhes:
                    url_partida = f"https://{regiao_match}.api.riotgames.com/lol/match/v5/matches/{match_id}"
                    status_partida, detalhes_brutos = requisitar_json(url_partida)
                    if status_partida == 200 and isinstance(detalhes_brutos, dict):
                        cache_detalhes[match_id] = reduzir_detalhes_partida(detalhes_brutos)
                    else:
                        cache_detalhes[match_id] = None
                    # Libera explicitamente a resposta completa assim que os metadados são extraídos.
                    detalhes_brutos = None
                detalhes = cache_detalhes[match_id]

                if not detalhes:
                    print(f"[PARTIDA NÃO VÁLIDA] {match_id} | detalhes indisponíveis; seguindo para a próxima.")
                    continue

                info = detalhes.get("info", {})
                duracao = info.get("gameDuration", 0) or 0
                fila = info.get("queueId")
                mapa = info.get("mapId")
                modo = info.get("gameMode")
                motivo = None
                if fila != queue_id:
                    motivo = f"fila={fila}, esperada={queue_id}"
                elif mapa != map_id:
                    motivo = f"mapa={mapa}, esperado={map_id}"
                elif modo != game_mode:
                    motivo = f"modo={modo}, esperado={game_mode}"
                elif duracao < duracao_minima_segundos:
                    motivo = f"duração={duracao / 60:.1f} min (mínimo 20 min)"
                elif not info.get("participants") or not info.get("teams"):
                    motivo = "detalhes incompletos (participantes/equipes ausentes)"

                if motivo:
                    print(f"[PARTIDA NÃO VÁLIDA] {match_id} | {motivo}.")
                    continue

                # Uma partida elegível que já entrou no dataset conta como conflito
                # para o jogador atual; não conta para sua meta, então busca outra.
                if match_id in partidas_unicas:
                    conflitos += 1
                    conflitos_jogador += 1
                    print(f"[CONFLITO] {match_id} | jogador atual: {nome}; já incluída por: {dono_partida[match_id]}. Buscando substituta.")
                    continue

                partidas_unicas[match_id] = detalhes
                dono_partida[match_id] = nome
                validas_jogador += 1
                print(f"[PARTIDA VÁLIDA] {match_id} | {duracao / 60:.1f} min | jogador {indice}/{len(jogadores)}: {validas_jogador}/{quantidade_partidas} | total únicas: {len(partidas_unicas)}")

            if len(ids) < tamanho_pagina:
                print(f"[FIM DO HISTÓRICO] A API retornou menos de {tamanho_pagina} IDs para {nome}.")
                break

        jogador["partidas_validas"] = validas_jogador
        print(f"Resumo {nome}: válidas únicas {validas_jogador}/{quantidade_partidas}; examinadas {analisadas_jogador}; conflitos {conflitos_jogador}.")
        if validas_jogador < quantidade_partidas:
            print("  AVISO: histórico esgotado ou ocorreram falhas antes de completar a meta.")

    print("\nResumo da extração:")
    print(f"  Jogadores processados: {len(jogadores)}")
    print(f"  Partidas únicas selecionadas: {len(partidas_unicas)}")
    print(f"  IDs de partidas examinados (inclui repetições entre históricos): {partidas_examindas}")
    print(f"  Conflitos: {conflitos}")
    return partidas_unicas


# ============================================================
# 5. TRANSFORMAR TIMELINE DIRETAMENTE EM SÉRIE POR EQUIPE
# ============================================================

def achatar_dicionario(dicionario: dict, prefixo: str = "") -> dict:
    resultado = {}
    if not isinstance(dicionario, dict):
        return resultado
    for chave, valor in dicionario.items():
        nome_coluna = f"{prefixo}.{chave}" if prefixo else str(chave)
        if isinstance(valor, dict):
            resultado.update(achatar_dicionario(valor, nome_coluna))
        elif isinstance(valor, list):
            resultado[nome_coluna] = json.dumps(valor, ensure_ascii=False)
        else:
            resultado[nome_coluna] = valor
    return resultado


def inteiro(valor):
    try:
        if valor is None:
            return None
        return int(float(valor))
    except (ValueError, TypeError):
        return None


def texto_evento(evento: dict, coluna: str) -> str:
    valor = evento.get(coluna)
    return "" if valor is None else str(valor).strip().upper()


def obter_vencedor(info: dict):
    for equipe in info.get("teams", []):
        if equipe.get("win") is True:
            return equipe.get("teamId")
    return None


def reduzir_detalhes_partida(detalhes: dict) -> dict:
    """Mantém somente metadados usados na validação e na agregação da timeline."""
    info = detalhes.get("info", {})
    participantes = [
        {
            "participantId": p.get("participantId"),
            "teamId": p.get("teamId"),
        }
        for p in info.get("participants", [])
        if isinstance(p, dict)
    ]
    equipes = [
        {"teamId": e.get("teamId"), "win": e.get("win")}
        for e in info.get("teams", [])
        if isinstance(e, dict)
    ]
    return {
        "info": {
            "queueId": info.get("queueId"),
            "mapId": info.get("mapId"),
            "gameMode": info.get("gameMode"),
            "gameDuration": info.get("gameDuration", 0),
            "gameStartTimestamp": info.get("gameStartTimestamp"),
            "gameCreation": info.get("gameCreation"),
            "participants": participantes,
            "teams": equipes,
        }
    }


def processar_timeline(match_id: str, detalhes: dict) -> list[dict] | None:
    """Retorna linhas da série temporal da partida sem materializar eventos brutos."""
    info = detalhes.get("info", {})
    vencedor = obter_vencedor(info)
    participantes = {}
    mapa_participantes = {}

    for participante in info.get("participants", []):
        pid = inteiro(participante.get("participantId"))
        team_id = inteiro(participante.get("teamId"))
        if pid is not None:
            participantes[pid] = participante
            if team_id in EQUIPES:
                mapa_participantes[pid] = team_id

    url_timeline = f"https://{regiao_match}.api.riotgames.com/lol/match/v5/matches/{match_id}/timeline"
    status, dados_timeline = requisitar_json(url_timeline)
    if status != 200 or not isinstance(dados_timeline, dict):
        print(f"[TIMELINE FALHOU] {match_id}; esta partida será omitida da planilha final.")
        return None

    info_timeline = dados_timeline.get("info", {})
    frames = info_timeline.get("frames", [])
    if not frames:
        print(f"[TIMELINE VAZIA] {match_id}; partida omitida.")
        return None

    # Estrutura agregada: (minuto normalizado, equipe) -> soma de atributos.
    agregados = defaultdict(lambda: defaultdict(float))
    chaves_com_valor = defaultdict(set)
    # Eventos por minuto/equipe; depois serão acumulados temporalmente.
    eventos_por_minuto = defaultdict(lambda: defaultdict(int))
    max_timestep = 0

    for indice_frame, frame in enumerate(frames):
        timestamp_frame = frame.get("timestamp", indice_frame * info_timeline.get("frameInterval", 60000)) or 0
        timestep_jogador = int(timestamp_frame // 60000)
        max_timestep = max(max_timestep, timestep_jogador)

        for pid_texto, dados_jogador in frame.get("participantFrames", {}).items():
            pid = inteiro(pid_texto)
            participante = participantes.get(pid, {})
            team_id = inteiro(participante.get("teamId"))
            if team_id not in EQUIPES:
                continue
            achatado = achatar_dicionario(dados_jogador)
            chave = (timestep_jogador, team_id)
            for coluna in COLUNAS_SOMA:
                if coluna not in achatado or achatado[coluna] is None:
                    continue
                try:
                    valor = float(achatado[coluna])
                    if math.isfinite(valor):
                        agregados[chave][coluna] += valor
                        chaves_com_valor[chave].add(coluna)
                except (ValueError, TypeError):
                    continue

        # Os eventos são atribuídos ao próximo minuto (ceil), como no notebook 2.
        for evento in frame.get("events", []):
            ts_evento = evento.get("timestamp", timestamp_frame) or 0
            timestep_evento = int(math.ceil(ts_evento / 60000))
            max_timestep = max(max_timestep, timestep_evento)
            tipo = texto_evento(evento, "type")
            killer_id = inteiro(evento.get("killerId"))
            killer_team = inteiro(evento.get("killerTeamId"))
            victim_id = inteiro(evento.get("victimId"))
            team_id_evento = inteiro(evento.get("teamId"))
            monster_type = texto_evento(evento, "monsterType")
            monster_subtype = texto_evento(evento, "monsterSubType")
            building_type = texto_evento(evento, "buildingType")

            if killer_team not in EQUIPES and killer_id is not None:
                killer_team = mapa_participantes.get(killer_id)
            victim_team = mapa_participantes.get(victim_id) if victim_id is not None else None

            contagens = {100: defaultdict(int), 200: defaultdict(int)}
            def adicionar(equipe, coluna, qtd=1):
                if equipe in EQUIPES:
                    contagens[equipe][coluna] += qtd

            if tipo == "CHAMPION_KILL":
                adicionar(killer_team, "kills")
                adicionar(victim_team, "deaths")
                assistentes = evento.get("assistingParticipantIds", []) or []
                if isinstance(assistentes, list):
                    for assistente in assistentes:
                        adicionar(mapa_participantes.get(inteiro(assistente)), "assists")

            if tipo == "ELITE_MONSTER_KILL":
                equipe_objetivo = killer_team if killer_team in EQUIPES else None
                if monster_type == "DRAGON":
                    adicionar(equipe_objetivo, "elder_dragons" if "ELDER" in monster_subtype else "dragons")
                elif monster_type in ("BARON_NASHOR", "BARON"):
                    adicionar(equipe_objetivo, "barons")
                elif monster_type in ("RIFTHERALD", "RIFT_HERALD"):
                    adicionar(equipe_objetivo, "rift_heralds")
                elif "ATAKHAN" in monster_type or "ATAKHAN" in monster_subtype:
                    adicionar(equipe_objetivo, "atakhan")

            if tipo == "BUILDING_KILL":
                equipe_objetivo = killer_team if killer_team in EQUIPES else None
                if equipe_objetivo not in EQUIPES and team_id_evento in EQUIPES:
                    equipe_objetivo = 300 - team_id_evento
                if "TOWER" in building_type:
                    adicionar(equipe_objetivo, "towers_destroyed")
                elif "INHIBITOR" in building_type:
                    adicionar(equipe_objetivo, "inhibitors_destroyed")

            for equipe in EQUIPES:
                for coluna, quantidade in contagens[equipe].items():
                    eventos_por_minuto[(timestep_evento, equipe)][coluna] += quantidade

    if not agregados:
        print(f"[SEM DADOS DE JOGADORES] {match_id}; partida omitida.")
        return None

    timestep_min = min(t for t, _ in agregados)
    timestep_max = max(max_timestep, max(t for t, _ in agregados))
    acumulados = {equipe: defaultdict(int) for equipe in EQUIPES}
    linhas = []

    # Grade completa por minuto e para ambas as equipes.
    for timestep in range(timestep_min, timestep_max + 1):
        for equipe in EQUIPES:
            for coluna, quantidade in eventos_por_minuto[(timestep, equipe)].items():
                acumulados[equipe][coluna] += quantidade
            chave = (timestep, equipe)
            linha = {
                "match_id": match_id,
                "timestep": timestep,
                "teamId": equipe,
                "vencedor": vencedor,
            }
            for coluna in COLUNAS_SOMA:
                linha[coluna] = agregados[chave].get(coluna, None) if coluna in chaves_com_valor[chave] else None
            for coluna in COLUNAS_EVENTOS:
                linha[coluna] = acumulados[equipe].get(coluna, 0)
            linhas.append(linha)

    # Diferenças azul - vermelho, iguais para as duas linhas do mesmo minuto.
    por_minuto = {t: {e: {} for e in EQUIPES} for t in range(timestep_min, timestep_max + 1)}
    for linha in linhas:
        por_minuto[linha["timestep"]][linha["teamId"]] = linha
    for timestep, equipes in por_minuto.items():
        azul, vermelho = equipes[100], equipes[200]
        for coluna in COLUNAS_SOMA + COLUNAS_EVENTOS:
            va, vb = azul.get(coluna, None), vermelho.get(coluna, None)
            try:
                diff = va - vb if va is not None and vb is not None else None
            except TypeError:
                diff = None
            azul[f"diff_{coluna}"] = diff
            vermelho[f"diff_{coluna}"] = diff

    return linhas


# ============================================================
# 6. GERAÇÃO DA PLANILHA FINAL — SEM ARQUIVO BRUTO INTERMEDIÁRIO
# ============================================================

def main():
    inicio_execucao = time.time()
    jogadores = selecionar_jogadores()
    partidas = coletar_partidas(jogadores)
    if not partidas:
        raise RuntimeError("Nenhuma partida válida foi selecionada; a planilha não foi criada.")

    # Só metadados compactos ficam em memória. A ordenação cronológica é feita antes
    # da consulta às timelines, permitindo gravar a saída na mesma ordem sem DataFrame.
    partidas_ordenadas = sorted(
        partidas.items(),
        key=lambda item: (
            item[1].get("info", {}).get("gameStartTimestamp")
            or item[1].get("info", {}).get("gameCreation")
            or 0,
            item[0],
        ),
    )

    print("\n" + "=" * 72)
    print("ETAPA 3/4 — TIMELINES E AGREGAÇÃO DIRETA POR EQUIPE")
    print("=" * 72)

    colunas_identificacao = ["match_id", "timestep", "teamId", "vencedor"]
    colunas_diff = [f"diff_{c}" for c in COLUNAS_SOMA + COLUNAS_EVENTOS]
    colunas_final = colunas_identificacao + COLUNAS_SOMA + COLUNAS_EVENTOS + colunas_diff

    # Write-only evita manter todas as células da planilha na RAM.
    workbook = Workbook(write_only=True)
    aba = workbook.create_sheet(title="Series_Temporais")
    aba.freeze_panes = "A2"
    aba.append(colunas_final)

    falhas_timeline = 0
    total = len(partidas_ordenadas)
    linhas_totais = 0

    for numero, (match_id, detalhes) in enumerate(partidas_ordenadas, start=1):
        print(f"[TIMELINE {numero}/{total} — {100 * numero / total:.1f}%] {match_id}")
        linhas = processar_timeline(match_id, detalhes)
        if linhas is None:
            falhas_timeline += 1
            continue

        # Cada partida é gravada imediatamente; as linhas deixam de ocupar RAM ao
        # avançar para a próxima partida. Não há DataFrame global nem ordenação final.
        for linha in linhas:
            aba.append([linha.get(coluna) for coluna in colunas_final])
        linhas_totais += len(linhas)
        print(f"  Série gravada: {len(linhas)} linhas | total gravado: {linhas_totais:,}")
        del linhas

    if linhas_totais == 0:
        raise RuntimeError("Nenhuma timeline foi processada com sucesso; a planilha não foi criada.")

    # 1 linha de cabeçalho + dados; Excel aceita no máximo 1.048.576 linhas por aba.
    if linhas_totais + 1 > 1_048_576:
        raise RuntimeError(
            f"A série final tem {linhas_totais:,} linhas e excede o limite do Excel por aba. "
            "Reduza quantidade_jogadores/quantidade_partidas ou exporte para CSV."
        )
    ultima_coluna = get_column_letter(len(colunas_final))
    aba.auto_filter.ref = f"A1:{ultima_coluna}{linhas_totais + 1}"
    workbook.save(nome_arquivo_saida)

    duracao = time.time() - inicio_execucao
    print("\n" + "=" * 72)
    print("CONCLUÍDO")
    print("=" * 72)
    print(f"Arquivo: {nome_arquivo_saida}")
    print(f"Partidas únicas selecionadas: {len(partidas)}")
    print(f"Timelines processadas com sucesso: {len(partidas) - falhas_timeline}")
    print(f"Timelines com falha: {falhas_timeline}")
    print(f"Linhas finais: {linhas_totais:,}")
    print(f"Colunas finais: {len(colunas_final)}")
    print(f"Tempo total: {duracao / 60:.1f} minutos")
    print(f"Planilha salva em: {nome_arquivo_saida}")


if __name__ == "__main__":
    main()
